from pathlib import Path

from openpyxl import load_workbook

from python_pipeline import excel_report


def test_matriz_marca_alerta_pdn_sin_qa(tmp_path):
    data = {
        "SOLO_QA.json": {"qa": {"presente": True, "subtipo": "carta", "ultima_vez": "2026-09-15"}},
        "AMBOS.json": {
            "qa": {"presente": True, "subtipo": "otro", "ultima_vez": "2026-09-15"},
            "pdn": {"presente": True, "subtipo": "otro", "ultima_vez": "2026-09-16"},
        },
        "SOLO_PDN.json": {"pdn": {"presente": True, "subtipo": "certificado", "ultima_vez": "2026-09-16"}},
        "ELIMINADO_DE_QA.json": {
            "qa": {"presente": False, "subtipo": "viejo", "ultima_vez": "2026-09-14"},
            "pdn": {"presente": True, "subtipo": "viejo", "ultima_vez": "2026-09-16"},
        },
    }

    out = tmp_path / "matriz.xlsx"
    excel_report.build_matriz_workbook(out, data, environments=("qa", "pdn", "dev"))

    wb = load_workbook(out)
    ws = wb["Matriz Ambientes"]
    rows = {row[0]: row for row in ws.iter_rows(min_row=2, values_only=True)}

    header = [c.value for c in ws[1]]
    idx_qa = header.index("QA")
    idx_pdn = header.index("PDN")
    idx_alerta = header.index("Alerta")

    assert rows["SOLO_QA.json"][idx_qa] == "Si"
    assert rows["SOLO_QA.json"][idx_alerta] in (None, "")

    assert rows["AMBOS.json"][idx_qa] == "Si"
    assert rows["AMBOS.json"][idx_pdn] == "Si"
    assert rows["AMBOS.json"][idx_alerta] in (None, "")

    assert rows["SOLO_PDN.json"][idx_qa] in (None, "")  # nunca visto en qa
    assert rows["SOLO_PDN.json"][idx_pdn] == "Si"
    assert rows["SOLO_PDN.json"][idx_alerta] == "EN PDN SIN QA"

    assert rows["ELIMINADO_DE_QA.json"][idx_qa] == "No"  # estuvo, ya no esta
    assert rows["ELIMINADO_DE_QA.json"][idx_pdn] == "Si"
    assert rows["ELIMINADO_DE_QA.json"][idx_alerta] == "EN PDN SIN QA"


def test_matriz_vacia_no_falla(tmp_path):
    out = tmp_path / "matriz.xlsx"
    excel_report.build_matriz_workbook(out, {}, environments=("qa", "pdn", "dev"))
    assert out.exists()

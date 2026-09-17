import csv
from pathlib import Path

from openpyxl import load_workbook

from python_pipeline import excel_report


def _write_group_csv(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow(["Nombre de los JSONs", "Nombre TA config", "Nombre del subtipo",
                          "Cantidad de archivos con ese subtipo", "Proceso", "Ruta JSON R3",
                          "Observaciones", "Transmisiones", "Tipo"])
        writer.writerow(["AAA.json", "ta1", "carta", "1", "p1", "s3://x/AAA", "", "Si", "R3 - Topics"])
        writer.writerow([])
        writer.writerow(["TOTAL_R3", "", "", "1", "", "", "", "", ""])


def test_build_maestro_workbook(tmp_path):
    group_csv = tmp_path / "grupo.csv"
    _write_group_csv(group_csv)

    out = tmp_path / "maestro.xlsx"
    excel_report.build_maestro_workbook(out, group_csv, resumen={"Ambiente": "qa", "Total de flujos conocidos (presentes)": 1})

    wb = load_workbook(out)
    assert wb.sheetnames == ["Resumen", "Reporte Agrupado"]

    ws_resumen = wb["Resumen"]
    valores = {row[0]: row[1] for row in ws_resumen.iter_rows(min_row=2, values_only=True)}
    assert valores["Ambiente"] == "qa"

    ws_grupo = wb["Reporte Agrupado"]
    rows = list(ws_grupo.iter_rows(values_only=True))
    assert rows[1][0] == "AAA.json"
    assert rows[1][2] == "carta"

import json
from pathlib import Path

from openpyxl import load_workbook

from python_pipeline import common, events_analysis


def _write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data), encoding="utf-8")


def test_classify_udz_tipo():
    assert common.classify_udz_tipo("s3://b/crudos/ns/X") == "crudos"
    assert common.classify_udz_tipo("s3://b/resultados/ns/X") == "resultados"
    assert common.classify_udz_tipo("s3://b/otra_cosa/ns/X") is None


def test_udz_flow_key_es_igual_para_crudos_y_resultados():
    crudos_path = "s3://b/crudos/ns_apoyo_corporativo/EVOLUCIONDIGITAL_SEGURODEACTIVO"
    resultados_path = "s3://b/resultados/ns_apoyo_corporativo/EVOLUCIONDIGITAL_SEGURODEACTIVO"
    key_crudos = common.udz_flow_key(crudos_path, "crudos")
    key_resultados = common.udz_flow_key(resultados_path, "resultados")
    assert key_crudos == key_resultados == "ns_apoyo_corporativo/EVOLUCIONDIGITAL_SEGURODEACTIVO"


def test_flujo_solo_crudos(tmp_path):
    events_dir = tmp_path / "events"
    _write_json(events_dir / "crudos" / "FLUJO_A.json", {
        "s3_path": "s3://b/crudos/ns_x/FLUJO_A",
    })

    flows = events_analysis.analyze(events_dir)

    assert len(flows) == 1
    assert flows[0]["flujo"] == "ns_x/FLUJO_A"
    assert flows[0]["tiene_crudos"] is True
    assert flows[0]["tiene_transmisiones"] is False
    assert flows[0]["observacion"] == "Solo Crudos"


def test_flujo_crudos_y_transmisiones(tmp_path):
    events_dir = tmp_path / "events"
    _write_json(events_dir / "crudos" / "FLUJO_B.json", {"s3_path": "s3://b/crudos/ns_x/FLUJO_B"})
    _write_json(events_dir / "resultados" / "FLUJO_B.json", {"s3_path": "s3://b/resultados/ns_x/FLUJO_B"})

    flows = events_analysis.analyze(events_dir)

    assert len(flows) == 1
    assert flows[0]["tiene_crudos"] is True
    assert flows[0]["tiene_transmisiones"] is True
    assert flows[0]["observacion"] == "Crudos + Transmisiones"


def test_flujo_solo_transmisiones_es_alerta(tmp_path):
    events_dir = tmp_path / "events"
    _write_json(events_dir / "resultados" / "FLUJO_C.json", {"s3_path": "s3://b/resultados/ns_x/FLUJO_C"})

    flows = events_analysis.analyze(events_dir)

    assert len(flows) == 1
    assert flows[0]["tiene_crudos"] is False
    assert flows[0]["observacion"].startswith("ALERTA")


def test_generate_escribe_excel_y_resumen(tmp_path):
    events_dir = tmp_path / "events"
    _write_json(events_dir / "crudos" / "A.json", {"s3_path": "s3://b/crudos/ns/A"})
    _write_json(events_dir / "crudos" / "B.json", {"s3_path": "s3://b/crudos/ns/B"})
    _write_json(events_dir / "resultados" / "B.json", {"s3_path": "s3://b/resultados/ns/B"})

    out_xlsx = tmp_path / "out.xlsx"
    summary = events_analysis.generate(events_dir, out_xlsx)

    assert summary["total_flujos"] == 2
    assert summary["solo_crudos"] == 1
    assert summary["crudos_y_transmisiones"] == 1
    assert summary["alertas_sin_crudos"] == 0

    assert out_xlsx.exists()
    ws = load_workbook(out_xlsx)["Eventos UDZ"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Flujo", "Tiene Crudos", "Tiene Transmisiones", "Observacion", "Archivo Crudos", "Archivo Transmisiones")
    assert len(rows) == 3  # header + 2 flujos
